"""
Builds the two PR2 decks from the user's PR2 template (research/presentation/PR2-TEMPLATE.pptx), keeping its
backgrounds, header bar, fonts and colors, and adding a red overlay on content slides so the text blends in.
  BUHAY-AH-PR2-Introduction-Methodology.pptx  (IMRaD: Introduction and Methodology)
  BUHAY-AH-PR2-Full-Research.pptx             (full research)
Run: python3 research/pipeline/build_pr2_decks.py
"""
import copy
import json
import os
import shutil

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TPL = os.path.join(ROOT, 'research', 'presentation', 'PR2-TEMPLATE.pptx')
FIG = os.path.join(ROOT, 'research', 'figures')
OUTS = [os.path.join(ROOT, 'docs', 'downloads'), os.path.join(ROOT, 'research')]
M = json.load(open(os.path.join(ROOT, 'research', 'outputs', 'metrics.json')))
DR = json.load(open(os.path.join(ROOT, 'research', 'outputs', 'design_results.json')))
X = json.load(open(os.path.join(ROOT, 'research', 'outputs', 'extension_results.json')))
B, R, P = M['building'], M['bridge'], M['pooled']
HY = next(v for k, v in P.items() if 'hybrid' in k)
ISO = P['Isolation Forest (CORA-style)']
XO = X['benchmark']['overall']
RED, DARK, WHITE = RGBColor(0x8B, 0x07, 0x07), RGBColor(0x5C, 0x03, 0x03), RGBColor(0xFF, 0xFF, 0xFF)
FONT = 'ITC Avant Garde Gothic'
TITLE = ('A Budget-to-Blueprint Design Generator, Standards-Based Quantity Take-Off, and Hybrid Machine Learning System for '
         'Start Quotations and Over-Budget Detection in Public Buildings, Bridges, Airports and Seaports')
RNS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def pct(x, d=1):
    return f'{x * 100:.{d}f}%'


# ------------------------------------------------------------------ slide cloning (keeps the template's art)
def clone(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    rmap = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith('/image'):
            rmap[rid] = new.part.relate_to(rel._target, rel.reltype)
    def fix(el):
        for node in el.iter():
            for a in ('embed', 'link', 'id'):
                k = '{%s}%s' % (RNS, a)
                if k in node.attrib and node.attrib[k] in rmap:
                    node.attrib[k] = rmap[node.attrib[k]]
        return el
    sbg = src.element.cSld.find(qn('p:bg'))
    if sbg is not None:
        new.element.cSld.insert(0, fix(copy.deepcopy(sbg)))
    tree = new.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn('p:nvGrpSpPr'), qn('p:grpSpPr')):
            continue
        tree.append(fix(copy.deepcopy(el)))
    return new


def shape(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def set_text(sh, text, size=None):
    p = sh.text_frame.paragraphs[0]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for extra in sh.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    if size:
        runs[0].font.size = Pt(size)


def overlay(slide):
    """semi-transparent red layer above the template background, below the header bar"""
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(20), Inches(11.25))
    r.fill.solid(); r.fill.fore_color.rgb = RED
    a = r.fill._xPr.find(qn('a:solidFill'))[0].makeelement(qn('a:alpha'), {'val': '72000'})
    r.fill._xPr.find(qn('a:solidFill'))[0].append(a)
    r.line.fill.background()
    r.name = 'Red overlay'
    tree = slide.shapes._spTree
    tree.remove(r._element)
    tree.insert(3, r._element)          # after nvGrpSpPr, grpSpPr and the background picture


def bullet(p, on=True, level=0):
    pPr = p._p.get_or_add_pPr()
    for t in ('a:buNone', 'a:buChar', 'a:buAutoNum'):
        for e in pPr.findall(qn(t)):
            pPr.remove(e)
    if on:
        pPr.set('marL', str(int(Inches(0.45 + 0.45 * level)))); pPr.set('indent', str(int(-Inches(0.35))))
        b = pPr.makeelement(qn('a:buChar'), {'char': '•'}); pPr.append(b)
    else:
        pPr.append(pPr.makeelement(qn('a:buNone'), {}))


def text(slide, x, y, w, h, items, size=30, color=WHITE, bullets=True, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=10):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for k, it in enumerate(items):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        parts = it if isinstance(it, list) else [(it, False)]
        for t, b in parts:
            r = p.add_run(); r.text = t
            f = r.font; f.size = Pt(size); f.bold = b; f.color.rgb = color; f.name = FONT
        bullet(p, bullets and not (isinstance(it, list) and it and it[0][0].startswith('§')))
        if isinstance(it, list) and it and it[0][0].startswith('§'):
            p.runs[0].text = p.runs[0].text[1:]
    return tb


def pic(slide, path, x, y, w, h, card=True):
    im = Image.open(path); r = min(w / im.size[0], h / im.size[1])
    iw, ih = im.size[0] * r, im.size[1] * r
    ax, ay = x + (w - iw) / 2, y + (h - ih) / 2
    if card:
        c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(ax - 0.12), Inches(ay - 0.12), Inches(iw + 0.24), Inches(ih + 0.24))
        c.fill.solid(); c.fill.fore_color.rgb = WHITE; c.line.fill.background(); c.adjustments[0] = 0.04
    slide.shapes.add_picture(path, Inches(ax), Inches(ay), Inches(iw), Inches(ih))


def table(slide, x, y, w, rows, widths, size=20, rowh=0.62):
    t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rowh * len(rows))).table
    for j, cw in enumerate(widths):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = DARK if i == 0 else (RGBColor(0x7A, 0x06, 0x06) if i % 2 else RGBColor(0x6A, 0x05, 0x05))
            c.text = str(v); p = c.text_frame.paragraphs[0]
            for r in p.runs:
                r.font.size = Pt(size); r.font.color.rgb = WHITE; r.font.bold = i == 0; r.font.name = FONT
            c.margin_left = c.margin_right = Inches(0.1)


def eqs(slide, keys, x, y, w, h):
    n = len(keys); gap = 0.25; each = (h - gap * (n - 1)) / n
    for k, key in enumerate(keys):
        pic(slide, os.path.join(FIG, 'eq', key + '.png'), x, y + k * (each + gap), w, each)


# ------------------------------------------------------------------ deck content
def build(full, name):
    prs = Presentation(TPL)
    T = list(prs.slides)
    ttl, content, qa, thanks = T[0], T[1], T[8], T[9]

    s = clone(prs, ttl)
    set_text(shape(s, 'TextBox 6'), 'BUHAY-AH', 86)
    set_text(shape(s, 'TextBox 9'), 'Enrique Miguel C. Buban · Lead Researcher')
    tb = shape(s, 'TextBox 6')
    text(s, tb.left / 914400 - 0.6, tb.top / 914400 + tb.height / 914400 + 0.05, tb.width / 914400 + 1.2, 1.6, [TITLE],
         size=17, bullets=False, align=PP_ALIGN.CENTER)

    def page(header, size=None):
        sl = clone(prs, content)
        hs = shape(sl, 'TextBox 6')
        hs.width = Inches(19.3)
        set_text(hs, header, 60 if len(header) <= 30 else 48 if len(header) <= 40 else 42)
        sl.shapes._spTree.remove(shape(sl, 'TextBox 7')._element)
        overlay(sl)
        return sl

    # ---------------- INTRODUCTION
    sl = page('INTRODUCTION')
    text(sl, 0.9, 2.3, 18.2, 8.6, [
        [('10–30% ', True), ('of public infrastructure investment is lost to corruption, mismanagement and inefficiency (CoST).', False)],
        [('32 / 100: ', True), ('Philippine score in the 2025 Corruption Perceptions Index, 120th of 182.', False)],
        [('₱545.64 billion ', True), ('in flood-control projects (2022–2025); about 20% went to 15 contractors.', False)],
        [('₱62.64 million ', True), ('overstatement in the Cebu City Medical Center: VAT at 12% instead of 5%.', False)],
        [('₱1.225 billion ', True), ('Cabagan–Santa Maria Bridge collapsed less than a month after completion (2025).', False)],
        'Outside reviewers cannot rebuild an estimate from the plans; padding hides in the price.'], size=30)
    sl = page('INTRODUCTION')
    text(sl, 0.9, 2.3, 10.3, 8.6, [
        'Schools, hospitals and evacuation centers protect communities in typhoon- and earthquake-prone Bicol.',
        'Airports, seaports, RoRo piers and jeepney, UV Express and bus terminals are the largest facilities a region builds.',
        'Machine learning learns what projects HAVE cost, not what they SHOULD cost.',
        'Communities often have only a need and a budget, not a plan.'], size=28)
    pic(sl, os.path.join(FIG, 'viz_building.png'), 11.6, 2.3, 3.6, 3.6)
    pic(sl, os.path.join(FIG, 'viz_bridge_truss.png'), 15.6, 2.3, 3.6, 3.6)
    pic(sl, os.path.join(FIG, 'airfield_runway.png'), 11.6, 6.4, 7.6, 1.8)
    pic(sl, os.path.join(FIG, 'port_roro.png'), 11.6, 8.5, 7.6, 2.2)
    sl = page('RESEARCH GAP AND RATIONALE', 50)
    text(sl, 0.9, 2.3, 18.2, 8.6, [
        [('Scope gap: ', True), ('cost screening has focused on roads, not buildings, bridges, airports and ports.', False)],
        [('Reference gap: ', True), ('data-driven models copy past padding; no open standards-based yardstick.', False)],
        [('Validation gap: ', True), ('detectors are never tested where the truly padded contracts are known.', False)],
        [('Design gap: ', True), ('no independent design and price for a community with only a need and a budget.', False)],
        [('Access gap: ', True), ('models stay in notebooks, not on the phones of auditors, officials and citizens.', False)],
        [('§Rationale: ', True), ('a fair price can be computed from the plan and the standards, then compared with the contract.', False)]], size=29)

    # ---------------- RRL/RRS
    sl = page('RRL/RRS')
    table(sl, 0.8, 2.3, 18.4, [
        ['Study', 'Method', 'Finding', 'Gap for this study'],
        ['Mahamid (2011)', 'Multiple regression', 'Early cost estimates', 'Roads only'],
        ['Elmousalami (2020)', 'AI cost-model review', 'Tree ensembles lead', 'Learns past prices'],
        ['Coffie & Cudjoe (2024)', 'XGBoost + SHAP', 'Explains cost overruns', 'No reference price'],
        ['Larocque et al. (2025)', 'Explainable ML', '11,646 road contracts', 'Roads only'],
        ['Liu et al. (2008)', 'Isolation Forest', 'Label-free anomalies', 'Misses ordinary padding'],
        ['Buban (2026), CORA', 'MLR + XGBoost + IF', 'R² 0.8723, MAPE 29.20%', 'Recall unknown'],
        ['GAO (2020)', 'Should-cost estimating', 'Bottom-up fair price', 'No open tool'],
        ['ICAO (2022); OCDI (2009)', 'Aerodrome / port standards', 'Runway, pier rules', 'Not priced'],
    ], [4.4, 4.6, 4.8, 4.6], size=21, rowh=0.86)

    # ---------------- FRAMEWORKS
    sl = page('THEORETICAL FRAMEWORK', 56)
    pic(sl, os.path.join(FIG, 'diagram_theoretical.png'), 0.8, 2.3, 18.4, 6.9)
    text(sl, 0.8, 9.6, 18.4, 1.2, ['Agency · Should-Cost · Optimization · Statistical Learning · Cooperative Game Theory'],
         size=24, bullets=False, align=PP_ALIGN.CENTER)
    sl = page('CONCEPTUAL FRAMEWORK PARADIGM', 46)
    pic(sl, os.path.join(FIG, 'diagram_conceptual.png'), 0.8, 2.3, 18.4, 8.5)

    # ---------------- SIGNIFICANCE
    sl = page('SIGNIFICANCE OF THE STUDY', 50)
    text(sl, 0.9, 2.3, 18.2, 8.6, [
        [('DPWH, DOH, DOTr, CAAP, PPA: ', True), ('cross-check programs of works and ABCs before bidding.', False)],
        [('COA, Ombudsman, ICI: ', True), ('rank contracts by risk; focus inspections.', False)],
        [('LGUs, schools, hospitals, cooperatives: ', True), ('a design and price from a need or a budget.', False)],
        [('Contractors: ', True), ('a public reference that rewards fair bids.', False)],
        [('Citizens and media: ', True), ('check any project on a phone.', False)],
        [('Academe: ', True), ('open pipeline, price book and benchmarks with known overpricing.', False)],
        [('SDGs: ', True), ('16.5, 16.6, 16.10 (anti-corruption) · 9.1, 11.5 (resilient infrastructure) · 12.7 (procurement) · 4.a (schools).', False)]], size=27)

    # ---------------- OBJECTIVES / SOP
    sl = page('OBJECTIVES OF THE STUDY / SOP', 46)
    text(sl, 0.9, 2.2, 18.2, 8.8, [
        [('§This study aimed to develop BUHAY-AH, a budget-to-blueprint design generator, standards-based quantity take-off and hybrid machine-learning system that produces start quotations and detects over-budgeted public buildings, bridges, airports and seaports in Bicol Region V. Specifically:', False)],
        '1. What plan variables can be extracted from plans and standards (physical, structural, location and year)?',
        '2. How can designs be generated from a requirement or a budget (module library, optimization, standards)?',
        '3. How accurate are the lean, standard and ceiling quotations against fair awarded costs (MAPE, R²)?',
        '4. How do MLR and XGBoost predict cost under LOOCV (R², RMSE, MAPE)?'], size=25, bullets=False)
    sl = page('OBJECTIVES OF THE STUDY / SOP', 46)
    text(sl, 0.9, 2.3, 18.2, 8.6, [
        '5. How well do the over-budget ratio, XGBoost residual, Isolation Forest and BUHAY-AH Risk Index detect over-budgeted contracts (precision, recall, F1, AUC)?',
        '6. How can the generator and engine extend to hospitals, airport, seaport and jeepney, UV Express and bus terminals, airfield pavements and port structures (templates, accuracy, detection)?',
        '7. How can the system be deployed as a mobile app, a download website and an explainable interface?'], size=28, bullets=False)

    sl = page('HYPOTHESIS (IF ANY)')
    text(sl, 0.9, 2.3, 18.2, 8.6, [
        [('Developmental study: ', True), ('expected outcomes instead of statistical hypotheses (DM 058 s. 2026 Research Plan).', False)],
        'The standards-based estimate will be closer to fair awarded costs than the data-driven models.',
        'A hybrid risk index will catch more padded contracts than an Isolation Forest alone.',
        'An exact budget-to-design search will run fast enough for a phone.',
        'The same take-off will extend to hospitals, terminals, runways and piers.'], size=29)

    # ---------------- METHODOLOGY
    sl = page('METHODOLOGY')
    text(sl, 0.9, 2.2, 18.2, 2.4, ['Quantitative developmental design with a simulation-based computational experiment (Mathematics and Computational Science): deterministic modeling, combinatorial optimization, statistical and machine learning, and Monte Carlo simulation.'], size=26, bullets=False)
    pic(sl, os.path.join(FIG, 'diagram_method.png'), 0.8, 4.8, 18.4, 5.8)
    sl = page('METHODOLOGY: DESIGN GENERATION', 44)
    text(sl, 0.9, 2.2, 8.6, 8.6, [
        f'{X["counts"]["spaces"]} space modules on DepEd, DOH, NBC, DSWD, IATA and AASHTO norms.',
        'Rows and floors by a greedy longest-first rule; stairs and exits per RA 9514.',
        'Budget-to-design: largest capacity whose quotation fits the budget.',
        'Exact enumeration + Pareto filter (cost is not monotonic).'], size=26)
    eqs(sl, ['pack', 'opt', 'pareto'], 9.9, 2.3, 9.3, 4.4)
    pic(sl, os.path.join(FIG, 'plan_school_gf.png'), 9.9, 7.1, 9.3, 3.6)
    sl = page('METHODOLOGY: QUANTITY TAKE-OFF', 44)
    text(sl, 0.9, 2.2, 8.6, 8.6, [
        'NSCP 2015, DPWH Vol. II/III, AASHTO LRFD, PD 1096, BP 344, RA 9514.',
        'DPWH pay items and 2025 Bicol unit costs (editable price book).',
        'OCM, profit and 5% VAT per D.O. 197 s. 2016 / D.O. 30 s. 2025.',
        'Three tiers: lean start quotation, standard fair estimate, ceiling.'], size=26)
    eqs(sl, ['footing', 'rebar', 'edc', 'abc', 'obr'], 9.9, 2.3, 9.3, 8.4)
    sl = page('METHODOLOGY: EXTENSION', 50)
    text(sl, 0.9, 2.2, 8.6, 8.6, [
        'Hospitals (DOH Level 1–3): medical gases, standby power, bed elevators.',
        'Airport and seaport terminals: curtain walls, baggage, curbside.',
        'Jeepney, UV Express and bus terminals: bays, PCCP, canopies.',
        'Runways, taxiways, aprons (ICAO, FAA); piers and wharves on piles (OCDI, PIANC).'], size=25)
    eqs(sl, ['bays', 'airfield', 'piles', 'bri_ext'], 9.9, 2.3, 9.3, 8.4)
    sl = page('METHODOLOGY: BENCHMARKS', 50)
    text(sl, 0.9, 2.2, 18.2, 4.8, [
        'SBD-300: 160 buildings + 140 bridges, 36 injected overpricing cases.',
        f'SBD-EXT: {XO["n"]} hospitals, terminals, airfields and ports, {XO["injected"]} injected cases.',
        'Hidden variation: unit prices ±7%, rebar ±10%, concrete ±6%, MEPF ±10%, piles ×0.95–1.35.',
        'Overpricing μ = 1.20–1.40 (moderate) and 1.40–2.10 (severe).'], size=27)
    eqs(sl, ['sim'], 3.0, 7.4, 14.0, 2.6)
    sl = page('METHODOLOGY: MODELS', 50)
    text(sl, 0.9, 2.2, 8.6, 8.6, [
        'MLR with VIF screening (interpretable baseline).',
        'XGBoost on log cost, randomized search, 5-fold CV.',
        'Isolation Forest on features, residual and OBR.',
        'Exact TreeSHAP explanations, also on the phone.'], size=26)
    eqs(sl, ['mlr', 'vif', 'xgb', 'iforest', 'shap'], 9.9, 2.3, 9.3, 8.4)
    sl = page('METHODOLOGY: VALIDATION', 50)
    text(sl, 0.9, 2.2, 8.6, 8.6, [
        'Leave-one-out cross-validation for cost models.',
        'RMSE, R² and MAPE for cost prediction.',
        'Precision, recall, F1 and ROC AUC for detection.',
        'BRI ≥ 30 flags; 20–30 review.'], size=26)
    eqs(sl, ['bri', 'bri_s', 'mape', 'prf'], 9.9, 2.3, 9.3, 8.4)
    sl = page('METHODOLOGY: APP & ETHICS', 50)
    pic(sl, os.path.join(FIG, 'diagram_architecture.png'), 0.8, 2.3, 18.4, 4.4)
    text(sl, 0.9, 7.1, 18.2, 3.8, [
        'One JavaScript engine for app, website and paper; offline PWA, Android APK, ZIP.',
        'No human participants, animals or hazardous materials; public and simulated data only.',
        'A flag calls for review by COA, the Ombudsman or the ICI; it is not proof of wrongdoing.'], size=26)

    # ---------------- RESULTS (full deck)
    if full:
        qb, qr = B['qto_vs_awarded_normal'], R['qto_vs_awarded_normal']
        sl = page('RESULTS: PLAN VARIABLES', 50)
        text(sl, 0.9, 2.3, 18.2, 8.6, [
            '9 building variables: GFA, footprint, storeys, span, floor height, type, roof, location, year.',
            '11 bridge variables: length, deck width and area, span, pier and abutment heights, superstructure, foundation, pile length, location, year.',
            'Extension: beds, peak-hour passengers, loading bays, runway and deck dimensions, water depth.',
            'Bridges varied far more than buildings in size and cost.'], size=28)
        sl = page('RESULTS: DESIGN GENERATION', 50)
        pic(sl, os.path.join(FIG, 'app_design_templates.png'), 0.8, 2.3, 11.0, 8.4)
        text(sl, 12.2, 2.3, 7.0, 8.4, [
            f'{X["counts"]["buildingTemplates"]} building, {X["counts"]["bridgeTemplates"]} bridge, {X["counts"]["airfieldTemplates"]} airfield, {X["counts"]["portTemplates"]} port designs.',
            f'{X["counts"]["checksPassed"]}/{X["counts"]["checks"]} space-standard checks passed.',
            f'Median search {DR["medianMs"]:.1f} ms.',
            f'Multi-storey adds up to {DR["maxGain"]} classrooms per budget.'], size=26)
        sl = page('RESULTS: BUDGET-TO-DESIGN', 50)
        pic(sl, os.path.join(FIG, 'frontier_school.png'), 0.8, 2.3, 11.0, 8.4)
        text(sl, 12.2, 2.3, 7.0, 8.4, ['₱10M → 4 classrooms (1 storey).', '₱25M → 9 classrooms (2 storeys).', '₱50M → 15 classrooms (4 storeys).',
             '₱150M → 185-bed hospital (4 storeys).', '₱40M → 12-bay bus terminal.'], size=26)
        sl = page('RESULTS: QUANTITY TAKE-OFF', 50)
        text(sl, 0.9, 2.3, 18.2, 8.6, [
            [('MAPE ', False), (f'{qb["MAPE"]:.2f}% ', True), ('buildings, ', False), (f'{qr["MAPE"]:.2f}% ', True), ('bridges; R² about 0.98.', False)],
            'Less than half the error of the best data-driven model.',
            f'Lean start quotation about {pct(B["lean_savings_vs_standard"])} (buildings) and {pct(R["lean_savings_vs_standard"])} (bridges) below standard, keeping code minimums.',
            f'Extension (SBD-EXT): MAPE {XO["accuracy"]["mape"]:.2f}%, R² {XO["accuracy"]["r2"]:.4f}; ceiling covers {XO["accuracy"]["ceilingCovers"]:.1f}% of fair awards.'], size=28)
        sl = page('RESULTS: COST MODELS')
        table(sl, 0.8, 2.4, 10.6, [['LOOCV', 'Buildings', 'Bridges'],
              ['MLR R²', f'{B["mlr_loocv"]["R2"]:.4f}', f'{R["mlr_loocv"]["R2"]:.4f}'], ['MLR MAPE', f'{B["mlr_loocv"]["MAPE"]:.2f}%', f'{R["mlr_loocv"]["MAPE"]:.2f}%'],
              ['XGBoost R²', f'{B["xgb_loocv"]["R2"]:.4f}', f'{R["xgb_loocv"]["R2"]:.4f}'], ['XGBoost MAPE', f'{B["xgb_loocv"]["MAPE"]:.2f}%', f'{R["xgb_loocv"]["MAPE"]:.2f}%']],
              [4.0, 3.3, 3.3], size=24, rowh=0.9)
        pic(sl, os.path.join(FIG, 'shap_bridge.png'), 11.8, 2.3, 7.4, 8.4)
        sl = page('RESULTS: DETECTION')
        pic(sl, os.path.join(FIG, 'roc_pooled.png'), 0.8, 2.3, 9.0, 8.4)
        text(sl, 10.3, 2.3, 8.9, 8.4, [
            f'Risk Index: recall {pct(HY["Recall"])}, F1 {HY["F1"]:.3f}, AUC {HY["AUC"]:.3f}.',
            f'Isolation Forest alone: recall {pct(ISO["Recall"])}.',
            f'Extension rule: recall {pct(XO["rule"]["recall"])}, precision {XO["rule"]["precision"]:.3f}.',
            'Most false alarms: unseen pile lengths (bridges, ports).'], size=27)
        sl = page('RESULTS: EXTENSION', 50)
        pic(sl, os.path.join(FIG, 'plan_hospital_l2.png'), 0.8, 2.3, 9.0, 3.4)
        pic(sl, os.path.join(FIG, 'airfield_runway.png'), 0.8, 6.0, 9.0, 2.2)
        pic(sl, os.path.join(FIG, 'port_roro.png'), 0.8, 8.5, 9.0, 2.3)
        pic(sl, os.path.join(FIG, 'ext_results.png'), 10.2, 2.3, 9.0, 8.5)
        sl = page('RESULTS: DEPLOYMENT', 50)
        pic(sl, os.path.join(FIG, 'app_design_budget.png'), 0.8, 2.3, 9.0, 8.4)
        pic(sl, os.path.join(FIG, 'app_check.png'), 10.2, 2.3, 9.0, 8.4)
        sl = page('CONCLUSION')
        text(sl, 0.9, 2.2, 18.2, 8.8, [
            'Plan variables from plans and standards are enough to price a project.',
            'Designs can be generated from a need or a budget, exactly and fast, and pass the standards.',
            'The standards-based estimate beats the data-driven models on accuracy.',
            'XGBoost is needed for bridges; regression suffices for buildings.',
            f'The Risk Index caught {pct(HY["Recall"])} of padded contracts vs. {pct(ISO["Recall"])} for Isolation Forest alone.',
            'The engine extends to hospitals, terminals, runways and piers.',
            'BUHAY-AH runs offline as an app, an Android package and a website.'], size=26)
        sl = page('RECOMMENDATIONS', 56)
        text(sl, 0.9, 2.2, 18.2, 8.8, [
            'Validate on actual DPWH, DOH, CAAP and PPA contracts.',
            'Review the design library with licensed architects and engineers.',
            'Add geotechnical data for bridge and port piles.',
            'Connect current DPWH and PSA price data.',
            'Retrain the models for the extension types; calibrate thresholds by class.',
            'Extend to flood control, water systems and roads with CORA.',
            'Pilot with oversight bodies; publish open contracting data.'], size=27)
        sl = page('REFERENCES')
        text(sl, 0.9, 2.2, 18.2, 8.8, [
            'Buban, E. M. C. (2026). CORA: Computational outlier recognition and analytics for road construction cost prediction.',
            'Chen, T., & Guestrin, C. (2016). XGBoost: A scalable tree boosting system. KDD, 785–794.',
            'Department of Public Works and Highways. (2016). D.O. No. 197, s. 2016: Approved budget for the contract.',
            'Federal Aviation Administration. (2021). Airport pavement design and evaluation (AC 150/5320-6G).',
            'International Civil Aviation Organization. (2022). Annex 14, Vol. I: Aerodrome design and operations (9th ed.).',
            'Liu, F. T., Ting, K. M., & Zhou, Z.-H. (2008). Isolation forest. IEEE ICDM, 413–422.',
            'Lundberg, S. M., & Lee, S.-I. (2017). A unified approach to interpreting model predictions. NeurIPS 30.',
            'OCDI. (2009). Technical standards and commentaries for port and harbour facilities in Japan.',
            'Shapley, L. S. (1953). A value for n-person games. Contributions to the Theory of Games II, 307–317.',
            'U.S. Government Accountability Office. (2020). Cost estimating and assessment guide (GAO-20-195G).'], size=19, bullets=False, space=6)

    clone(prs, qa)
    clone(prs, thanks)
    # drop the ten original template slides
    lst = prs.slides._sldIdLst
    for sid in list(lst)[:len(T)]:
        prs.part.drop_rel(sid.get(qn('r:id'))); lst.remove(sid)
    for d in OUTS:
        prs.save(os.path.join(d, name))
    return len(prs.slides)


if __name__ == '__main__':
    print(build(False, 'BUHAY-AH-PR2-Introduction-Methodology.pptx'), build(True, 'BUHAY-AH-PR2-Full-Research.pptx'))
